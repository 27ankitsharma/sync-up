#!/usr/bin/env python3
"""Build syllabus JSON from the author-friendly syllabus.xlsx workbook.

The workbook source of truth has four sheets:
Tracks, Subjects, Modules, and Topics. The frontend still consumes generated
JSON from public/syllabus.json and src/data/syllabus.json.
"""

from __future__ import annotations

import json
import re
import shutil
import sys
from collections import Counter, OrderedDict, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook, load_workbook
except ModuleNotFoundError as exc:
    print(
        "Error: openpyxl is required to read/write syllabus.xlsx. "
        "Install it in your active Python environment or run this script with a Python that has openpyxl.",
        file=sys.stderr,
    )
    raise SystemExit(1) from exc


INPUT_FILE = Path("syllabus.xlsx")
LEGACY_BACKUP_FILE = Path("syllabus.legacy-backup.xlsx")
OUTPUT_FILE = Path("public/syllabus.json")
SRC_OUTPUT_FILE = Path("src/data/syllabus.json")
REPORT_FILE = Path("syllabus_migration_report.md")
BUILD_REPORT_FILE = Path("syllabus_build_report.md")
PREVIOUS_JSON_FILE = SRC_OUTPUT_FILE

LEGACY_SHEET = "Syllabus"
REQUIRED_SHEETS = ["Tracks", "Subjects", "Modules", "Topics"]

LEGACY_COLUMNS = [
    "Track",
    "Subject",
    "Module",
    "Topic",
    "Layer",
    "Difficulty",
    "Roles",
    "Status",
    "Priority",
    "Is_Radar",
    "Radar_Week",
    "Summary",
    "Why_It_Matters",
]

TRACK_COLUMNS = [
    "Track_ID",
    "Track",
    "Order",
    "Summary",
    "Course_ID(s)",
    "Knowledge_Layer",
    "Why_It_Matters",
    "Resources",
    "Is_Radar",
    "Radar_Start_Date",
    "Radar_End_Date",
    "Difficulty",
    "Learning_Time",
    "Lens_Relevance",
    "Content_Status",
]

SUBJECT_COLUMNS = [
    "Subject_ID",
    "Track_ID",
    "Subject",
    "Order",
    "Summary",
    "Course_ID(s)",
    "Knowledge_Layer",
    "Why_It_Matters",
    "Resources",
    "Is_Radar",
    "Radar_Start_Date",
    "Radar_End_Date",
    "Difficulty",
    "Learning_Time",
    "Lens_Relevance",
    "Content_Status",
]

MODULE_COLUMNS = [
    "Module_ID",
    "Track_ID",
    "Subject_ID",
    "Module",
    "Order",
    "Summary",
    "Course_ID(s)",
    "Knowledge_Layer",
    "Why_It_Matters",
    "Resources",
    "Is_Radar",
    "Radar_Start_Date",
    "Radar_End_Date",
    "Difficulty",
    "Learning_Time",
    "Lens_Relevance",
    "Content_Status",
]

TOPIC_COLUMNS = [
    "Topic_ID",
    "Track_ID",
    "Subject_ID",
    "Module_ID",
    "Topic",
    "Knowledge_Layer",
    "Difficulty",
    "Learning_Time",
    "Lens_Relevance",
    "Summary",
    "Why_It_Matters",
    "Resources",
    "Content_Status",
    "Course_Status",
    "Diagnostic_Status",
    "Is_Radar",
    "Radar_Start_Date",
    "Radar_End_Date",
    "Radar_Classification",
]

VALID_DIFFICULTIES = {"Beginner", "Intermediate", "Advanced"}
VALID_CONTENT_STATUSES = {"Draft", "Published", "Archived"}
VALID_AVAILABILITY_STATUSES = {"yes", "no", "WIP"}
AVAILABILITY_ALIASES = {"yes": "yes", "no": "no", "wip": "WIP"}
VALID_RELEVANCE = {"Must", "Good", "Optional"}
RADAR_CLASSIFICATION_ALIASES = {
    "new topic": "new_topic_candidate",
    "new_topic_candidate": "new_topic_candidate",
    "topic update": "existing_topic_update",
    "existing_topic_update": "existing_topic_update",
    "fyi": "fyi",
}
VALID_KNOWLEDGE_LAYERS = {
    "Foundations",
    "Models & Architectures",
    "Techniques & Practices",
    "Systems & Applications",
    "Frontiers & Emerging",
}
STATUS_TO_JSON = {"Draft": "draft", "Published": "published", "Archived": "archived"}
LEGACY_STATUS_MAP = {
    "draft": "Draft",
    "published": "Published",
    "coming_soon": "Draft",
    "coming-soon": "Draft",
    "archived": "Archived",
}
LEGACY_PRIORITY_TO_RELEVANCE = {"high": "Must", "medium": "Good", "low": "Optional"}
RELEVANCE_TO_PRIORITY = {"Must": "high", "Good": "medium", "Optional": "low"}
ISO_WEEK_RE = re.compile(r"^(\d{4})-W(\d{2})$")


@dataclass
class BuildIssueLog:
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    manual_review: list[str] = field(default_factory=list)
    migrated: bool = False
    migrated_records: dict[str, int] = field(default_factory=dict)

    def error(self, message: str) -> None:
        self.errors.append(message)

    def warn(self, message: str) -> None:
        self.warnings.append(message)

    def review(self, message: str) -> None:
        self.manual_review.append(message)


def slugify(value: str) -> str:
    slug = value.strip().lower()
    slug = re.sub(r"[^a-z0-9\s-]", "", slug)
    slug = re.sub(r"[\s_-]+", "-", slug)
    return slug.strip("-")


def cell_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value).strip()


def load_workbook_rows(path: Path, sheet_name: str, expected_columns: list[str]) -> list[dict[str, str]]:
    wb = load_workbook(path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Missing required sheet: {sheet_name}")

    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"Sheet {sheet_name} is empty")

    headers = [cell_value(value) for value in rows[0]][: len(expected_columns)]
    if headers != expected_columns:
        raise ValueError(
            f"Invalid columns in {sheet_name}.\n"
            f"Expected: {expected_columns}\n"
            f"Found:    {headers}"
        )

    result: list[dict[str, str]] = []
    for raw_row in rows[1:]:
        values = [cell_value(value) for value in raw_row]
        padded = values + [""] * (len(expected_columns) - len(values))
        if not any(value.strip() for value in padded):
            continue
        result.append(dict(zip(expected_columns, padded[: len(expected_columns)])))
    return result


def has_target_workbook(path: Path) -> bool:
    wb = load_workbook(path, read_only=True, data_only=True)
    return all(sheet in wb.sheetnames for sheet in REQUIRED_SHEETS)


def migrate_legacy_workbook(path: Path, log: BuildIssueLog) -> None:
    wb = load_workbook(path, data_only=True)
    if LEGACY_SHEET not in wb.sheetnames:
        raise ValueError(f"Workbook must contain {REQUIRED_SHEETS} or legacy sheet {LEGACY_SHEET}")

    ws = wb[LEGACY_SHEET]
    rows = list(ws.iter_rows(values_only=True))
    headers = [cell_value(value) for value in rows[0]][: len(LEGACY_COLUMNS)]
    if headers != LEGACY_COLUMNS:
        raise ValueError(
            "Legacy workbook columns are not recognized.\n"
            f"Expected: {LEGACY_COLUMNS}\n"
            f"Found:    {headers}"
        )

    legacy_rows: list[dict[str, str]] = []
    for raw_row in rows[1:]:
        values = [cell_value(value) for value in raw_row]
        padded = values + [""] * (len(LEGACY_COLUMNS) - len(values))
        if not any(value.strip() for value in padded):
            continue
        legacy_rows.append(dict(zip(LEGACY_COLUMNS, padded[: len(LEGACY_COLUMNS)])))

    if not LEGACY_BACKUP_FILE.exists():
        shutil.copy2(path, LEGACY_BACKUP_FILE)
    else:
        log.warn(f"Legacy backup already exists and was not overwritten: {LEGACY_BACKUP_FILE}")

    track_rows: OrderedDict[str, dict[str, Any]] = OrderedDict()
    subject_rows: OrderedDict[tuple[str, str], dict[str, Any]] = OrderedDict()
    module_rows: OrderedDict[tuple[str, str], dict[str, Any]] = OrderedDict()
    topic_rows: list[dict[str, Any]] = []
    topic_ids: Counter[str] = Counter()

    for index, row in enumerate(legacy_rows, start=2):
        track = row["Track"].strip()
        subject = row["Subject"].strip()
        module = row["Module"].strip()
        topic = row["Topic"].strip()
        if not all([track, subject, module, topic]):
            log.review(f"Legacy row {index}: missing hierarchy field; row migrated with available values")

        track_id = slugify(track)
        subject_id = slugify(subject)
        module_id = slugify(module)
        topic_id = slugify(topic)
        topic_ids[topic_id] += 1
        if topic_ids[topic_id] > 1:
            topic_id = f"{topic_id}-{topic_ids[topic_id]}"
            log.review(f"Legacy row {index}: duplicate generated Topic_ID; assigned {topic_id}")
        else:
            log.review(f"Legacy row {index}: Topic_ID generated from Topic because no permanent ID existed")

        track_rows.setdefault(
            track_id,
            {
                "Track_ID": track_id,
                "Track": track,
                "Order": len(track_rows) + 1,
                "Summary": "",
                "Course_ID(s)": "",
                "Knowledge_Layer": "",
                "Why_It_Matters": "",
                "Resources": "",
                "Is_Radar": "",
                "Radar_Start_Date": "",
                "Radar_End_Date": "",
                "Difficulty": "",
                "Learning_Time": "",
                "Lens_Relevance": "",
                "Content_Status": "Published",
            },
        )
        subject_rows.setdefault(
            (track_id, subject_id),
            {
                "Subject_ID": subject_id,
                "Track_ID": track_id,
                "Subject": subject,
                "Order": sum(1 for item in subject_rows if item[0] == track_id) + 1,
                "Summary": "",
                "Course_ID(s)": "",
                "Knowledge_Layer": "",
                "Why_It_Matters": "",
                "Resources": "",
                "Is_Radar": "",
                "Radar_Start_Date": "",
                "Radar_End_Date": "",
                "Difficulty": "",
                "Learning_Time": "",
                "Lens_Relevance": "",
                "Content_Status": "Published",
            },
        )
        module_rows.setdefault(
            (subject_id, module_id),
            {
                "Module_ID": module_id,
                "Track_ID": track_id,
                "Subject_ID": subject_id,
                "Module": module,
                "Order": sum(1 for item in module_rows if item[0] == subject_id) + 1,
                "Summary": "",
                "Course_ID(s)": "",
                "Knowledge_Layer": row["Layer"].strip(),
                "Why_It_Matters": "",
                "Resources": "",
                "Is_Radar": "",
                "Radar_Start_Date": "",
                "Radar_End_Date": "",
                "Difficulty": "",
                "Learning_Time": "",
                "Lens_Relevance": "",
                "Content_Status": "Published",
            },
        )

        radar_start_date = ""
        radar_week = row["Radar_Week"].strip()
        if radar_week:
            try:
                radar_start_date = iso_week_to_date(radar_week).isoformat()
            except ValueError:
                log.review(f"Legacy row {index}: invalid Radar_Week requires manual review: {radar_week}")

        topic_rows.append(
            {
                "Topic_ID": topic_id,
                "Track_ID": track_id,
                "Subject_ID": subject_id,
                "Module_ID": module_id,
                "Topic": topic,
                "Knowledge_Layer": row["Layer"].strip(),
                "Difficulty": row["Difficulty"].strip(),
                "Learning_Time": "",
                "Lens_Relevance": legacy_lens_relevance(row),
                "Summary": row["Summary"].strip(),
                "Why_It_Matters": row["Why_It_Matters"].strip(),
                "Resources": "",
                "Content_Status": LEGACY_STATUS_MAP.get(row["Status"].strip().lower(), "Draft"),
                "Course_Status": "yes" if LEGACY_STATUS_MAP.get(row["Status"].strip().lower(), "Draft") == "Published" else "no",
                "Diagnostic_Status": "yes" if LEGACY_STATUS_MAP.get(row["Status"].strip().lower(), "Draft") == "Published" else "no",
                "Is_Radar": row["Is_Radar"].strip(),
                "Radar_Start_Date": radar_start_date,
                "Radar_End_Date": "",
            }
        )

    migrated = Workbook()
    migrated.remove(migrated.active)
    append_sheet(migrated, "Tracks", TRACK_COLUMNS, list(track_rows.values()))
    append_sheet(migrated, "Subjects", SUBJECT_COLUMNS, list(subject_rows.values()))
    append_sheet(migrated, "Modules", MODULE_COLUMNS, list(module_rows.values()))
    append_sheet(migrated, "Topics", TOPIC_COLUMNS, topic_rows)
    migrated.save(path)

    log.migrated = True
    log.migrated_records = {
        "Tracks": len(track_rows),
        "Subjects": len(subject_rows),
        "Modules": len(module_rows),
        "Topics": len(topic_rows),
    }


def append_sheet(wb: Workbook, title: str, columns: list[str], rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title)
    ws.append(columns)
    for row in rows:
        ws.append([row.get(column, "") for column in columns])
    for cell in ws[1]:
        cell.style = "Headline 4"
    ws.freeze_panes = "A2"


def legacy_lens_relevance(row: dict[str, str]) -> str:
    roles = [role.strip() for role in row["Roles"].split(",") if role.strip()]
    category = LEGACY_PRIORITY_TO_RELEVANCE.get(row["Priority"].strip().lower(), "Optional")
    return "; ".join(f"{role}: {category}" for role in roles)


def require_value(row: dict[str, str], column: str, label: str, log: BuildIssueLog) -> str:
    value = row.get(column, "").strip()
    if not value:
        log.error(f"{label}: {column} is required")
    return value


def parse_order(value: str, label: str, log: BuildIssueLog) -> int:
    if not value:
        return 999999
    try:
        return int(float(value))
    except ValueError:
        log.error(f"{label}: Order must be a number")
        return 999999


def parse_bool(value: str, label: str, log: BuildIssueLog) -> bool:
    normalized = value.strip().lower()
    if normalized in {"true", "1", "yes", "y"}:
        return True
    if normalized in {"false", "0", "no", "n", ""}:
        return False
    log.error(f"{label}: Is_Radar must be TRUE or FALSE")
    return False


def parse_date(value: str, label: str, column: str, log: BuildIssueLog) -> str | None:
    if not value:
        return None
    try:
        return datetime.strptime(value, "%Y-%m-%d").date().isoformat()
    except ValueError:
        log.error(f"{label}: {column} must use YYYY-MM-DD")
        return None


def iso_week_to_date(iso_week: str) -> date:
    match = ISO_WEEK_RE.match(iso_week)
    if not match:
        raise ValueError(f"Invalid ISO week: {iso_week}")
    year = int(match.group(1))
    week = int(match.group(2))
    return date.fromisocalendar(year, week, 1)


def date_to_iso_week(value: str | None) -> str | None:
    if not value:
        return None
    parsed = datetime.strptime(value, "%Y-%m-%d").date()
    year, week, _weekday = parsed.isocalendar()
    return f"{year}-W{week:02d}"


def parse_content_status(value: str, label: str, log: BuildIssueLog) -> str:
    normalized = value.strip()
    if normalized not in VALID_CONTENT_STATUSES:
        log.error(f"{label}: Content_Status must be one of {sorted(VALID_CONTENT_STATUSES)}")
        return "Draft"
    return normalized


def parse_availability_status(value: str, column: str, label: str, log: BuildIssueLog) -> str:
    normalized = value.strip()
    if not normalized:
        return "no"
    mapped = AVAILABILITY_ALIASES.get(normalized.lower())
    if mapped is None:
        log.error(f"{label}: {column} must be one of yes, no, WIP")
        return "no"
    return mapped


def parse_radar_classification(value: str, is_radar: bool, label: str, log: BuildIssueLog) -> str | None:
    normalized = value.strip().lower()
    if not normalized:
        if is_radar:
            log.error(f"{label}: Radar_Classification is required when Is_Radar is TRUE")
        return None
    classification = RADAR_CLASSIFICATION_ALIASES.get(normalized)
    if classification is None:
        log.error(f"{label}: Radar_Classification must be New Topic, Topic Update, or FYI")
    return classification


def ensure_topic_authoring_columns(path: Path) -> None:
    """Add newer topic authoring columns when the workbook predates them.

    Existing availability rows are initialized to yes so current learning experiences stay reachable.
    Existing Radar rows default to Topic Update; later Radar authoring must classify new entries explicitly.
    """
    wb = load_workbook(path)
    if "Topics" not in wb.sheetnames:
        return

    ws = wb["Topics"]
    changed = False
    for column, after in (("Course_Status", "Content_Status"), ("Diagnostic_Status", "Course_Status")):
        headers = [cell.value for cell in ws[1]]
        if column in headers:
            continue
        insert_at = headers.index(after) + 2 if after in headers else len(headers) + 1
        ws.insert_cols(insert_at)
        ws.cell(1, insert_at).value = column
        for row_index in range(2, ws.max_row + 1):
            if any(cell_value(ws.cell(row_index, col_index).value) for col_index in range(1, ws.max_column + 1)):
                ws.cell(row_index, insert_at).value = "yes"
        changed = True

    headers = [cell.value for cell in ws[1]]
    if "Radar_Classification" not in headers:
        after = "Radar_End_Date"
        insert_at = headers.index(after) + 2 if after in headers else len(headers) + 1
        ws.insert_cols(insert_at)
        ws.cell(1, insert_at).value = "Radar_Classification"
        headers = [cell.value for cell in ws[1]]
        radar_column = headers.index("Is_Radar") + 1
        for row_index in range(2, ws.max_row + 1):
            is_radar = cell_value(ws.cell(row_index, radar_column).value).strip().lower()
            if is_radar in {"true", "1", "yes", "y"}:
                ws.cell(row_index, insert_at).value = "Topic Update"
        changed = True

    if changed:
        wb.save(path)
    wb.close()


def parse_difficulty(value: str, label: str, log: BuildIssueLog) -> str:
    if value not in VALID_DIFFICULTIES:
        log.error(f"{label}: Difficulty must be one of {sorted(VALID_DIFFICULTIES)}")
        return "Beginner"
    return value


def parse_knowledge_layer(value: str, label: str, log: BuildIssueLog) -> str:
    normalized = value.strip()
    if normalized not in VALID_KNOWLEDGE_LAYERS:
        log.error(f"{label}: Knowledge_Layer must be one of {sorted(VALID_KNOWLEDGE_LAYERS)}")
        return normalized
    return normalized


def parse_learning_time(value: str, label: str, log: BuildIssueLog) -> float | None:
    if not value:
        return None
    normalized = value.strip().lower()
    try:
        if normalized.endswith("min"):
            return round(float(normalized[:-3].strip()) / 60, 2)
        if normalized.endswith("m"):
            return round(float(normalized[:-1].strip()) / 60, 2)
        if normalized.endswith("hours"):
            return round(float(normalized[:-5].strip()), 2)
        if normalized.endswith("hour"):
            return round(float(normalized[:-4].strip()), 2)
        if normalized.endswith("h"):
            return round(float(normalized[:-1].strip()), 2)
        return round(float(normalized), 2)
    except ValueError:
        log.error(f"{label}: Learning_Time must be a number of hours or a value like 90m / 1.5h")
        return None


def parse_course_ids(value: str) -> list[str]:
    if not value:
        return []
    return [item.strip() for item in re.split(r"[,;\n]", value) if item.strip()]


def parse_resources(value: str, label: str, log: BuildIssueLog) -> list[dict[str, str]]:
    if not value:
        return []
    stripped = value.strip()
    if stripped.startswith("[") or stripped.startswith("{"):
        try:
            parsed = json.loads(stripped)
        except json.JSONDecodeError as exc:
            log.error(f"{label}: Resources JSON is invalid: {exc}")
            return []
        items = parsed if isinstance(parsed, list) else [parsed]
        return normalize_resources(items, label, log)

    resources: list[dict[str, str]] = []
    for index, part in enumerate(re.split(r"\n|;;", stripped), start=1):
        item = part.strip()
        if not item:
            continue
        pieces = [piece.strip() for piece in item.split("|")]
        if len(pieces) != 3:
            log.error(f"{label}: Resource {index} must use Type | Title | URL")
            continue
        resources.append({"type": pieces[0], "title": pieces[1], "url": pieces[2]})
    return resources


def normalize_resources(items: list[Any], label: str, log: BuildIssueLog) -> list[dict[str, str]]:
    resources: list[dict[str, str]] = []
    for index, item in enumerate(items, start=1):
        if not isinstance(item, dict):
            log.error(f"{label}: Resource {index} must be an object")
            continue
        resource = {
            "type": str(item.get("type", "")).strip(),
            "title": str(item.get("title", "")).strip(),
            "url": str(item.get("url", "")).strip(),
        }
        if not all(resource.values()):
            log.error(f"{label}: Resource {index} requires type, title, and url")
            continue
        resources.append(resource)
    return resources


def parse_lens_relevance(value: str, label: str, log: BuildIssueLog) -> dict[str, str]:
    if not value:
        return {}
    stripped = value.strip()
    if stripped.startswith("{"):
        try:
            parsed = json.loads(stripped)
        except json.JSONDecodeError as exc:
            log.error(f"{label}: Lens_Relevance JSON is invalid: {exc}")
            return {}
        return normalize_lens_relevance(parsed, label, log)

    result: dict[str, str] = {}
    for part in re.split(r";|\n", stripped):
        item = part.strip()
        if not item:
            continue
        if ":" in item:
            role, category = item.split(":", 1)
        elif "=" in item:
            role, category = item.split("=", 1)
        else:
            log.error(f"{label}: Lens_Relevance item must use Role: Must/Good/Optional")
            continue
        role = role.strip()
        category = normalize_relevance_category(category.strip())
        if not role or category not in VALID_RELEVANCE:
            log.error(f"{label}: invalid Lens_Relevance item {item!r}")
            continue
        result[role] = category
    return result


def normalize_lens_relevance(parsed: Any, label: str, log: BuildIssueLog) -> dict[str, str]:
    if not isinstance(parsed, dict):
        log.error(f"{label}: Lens_Relevance JSON must be an object")
        return {}
    result: dict[str, str] = {}
    for role, raw_category in parsed.items():
        category = raw_category.get("category") if isinstance(raw_category, dict) else raw_category
        normalized = normalize_relevance_category(str(category).strip())
        if normalized not in VALID_RELEVANCE:
            log.error(f"{label}: Lens_Relevance for {role} must be Must, Good, or Optional")
            continue
        result[str(role).strip()] = normalized
    return result


def normalize_relevance_category(value: str) -> str:
    normalized = value.strip().lower()
    if normalized in {"must", "must learn", "high"}:
        return "Must"
    if normalized in {"good", "good to have", "medium"}:
        return "Good"
    if normalized in {"optional", "low"}:
        return "Optional"
    return value


def validate_unique(rows: list[dict[str, str]], column: str, sheet: str, log: BuildIssueLog) -> None:
    seen: dict[str, int] = {}
    for index, row in enumerate(rows, start=2):
        value = row.get(column, "").strip()
        if not value:
            log.error(f"{sheet} row {index}: {column} is required")
            continue
        if value in seen:
            log.error(f"{sheet} row {index}: duplicate {column} {value!r}; first seen on row {seen[value]}")
        seen[value] = index


def build_syllabus(
    track_rows: list[dict[str, str]],
    subject_rows: list[dict[str, str]],
    module_rows: list[dict[str, str]],
    topic_rows: list[dict[str, str]],
    log: BuildIssueLog,
) -> dict[str, Any]:
    validate_unique(track_rows, "Track_ID", "Tracks", log)
    validate_unique(subject_rows, "Subject_ID", "Subjects", log)
    validate_unique(module_rows, "Module_ID", "Modules", log)
    validate_unique(topic_rows, "Topic_ID", "Topics", log)

    track_by_id: dict[str, dict[str, Any]] = {}
    track_by_title: dict[str, dict[str, Any]] = {}
    for index, row in enumerate(track_rows, start=2):
        label = f"Tracks row {index}"
        track_id = require_value(row, "Track_ID", label, log)
        title = require_value(row, "Track", label, log)
        parsed = parse_hierarchy_row(row, "track", track_id, title, label, log)
        track_by_id[track_id] = parsed
        track_by_title[title] = parsed

    subject_by_id: dict[str, dict[str, Any]] = {}
    for index, row in enumerate(subject_rows, start=2):
        label = f"Subjects row {index}"
        subject_id = require_value(row, "Subject_ID", label, log)
        track_id = require_value(row, "Track_ID", label, log)
        title = require_value(row, "Subject", label, log)
        if track_id not in track_by_id:
            log.error(f"{label}: Track_ID {track_id!r} does not exist in Tracks sheet")
        parsed = parse_hierarchy_row(row, "subject", subject_id, title, label, log)
        parsed["track_id"] = track_id
        subject_by_id[subject_id] = parsed

    module_by_id: dict[str, dict[str, Any]] = {}
    for index, row in enumerate(module_rows, start=2):
        label = f"Modules row {index}"
        module_id = require_value(row, "Module_ID", label, log)
        track_id = require_value(row, "Track_ID", label, log)
        subject_id = require_value(row, "Subject_ID", label, log)
        title = require_value(row, "Module", label, log)
        if track_id not in track_by_id:
            log.error(f"{label}: Track_ID {track_id!r} does not exist in Tracks sheet")
        if subject_id not in subject_by_id:
            log.error(f"{label}: Subject_ID {subject_id!r} does not exist in Subjects sheet")
        elif subject_by_id[subject_id].get("track_id") != track_id:
            log.error(
                f"{label}: Track_ID {track_id!r} does not match Subject_ID {subject_id!r} "
                f"(expected {subject_by_id[subject_id].get('track_id')!r})"
            )
        parsed = parse_hierarchy_row(row, "module", module_id, title, label, log)
        parsed["track_id"] = track_id
        parsed["subject_id"] = subject_id
        module_by_id[module_id] = parsed

    previous_locations = load_previous_topic_locations()
    topics_by_module: dict[str, list[dict[str, Any]]] = defaultdict(list)
    topic_name_by_module: Counter[tuple[str, str]] = Counter()

    for index, row in enumerate(topic_rows, start=2):
        label = f"Topics row {index}"
        topic_id = require_value(row, "Topic_ID", label, log)
        track_id = require_value(row, "Track_ID", label, log)
        subject_id = require_value(row, "Subject_ID", label, log)
        module_id = require_value(row, "Module_ID", label, log)
        topic_title = require_value(row, "Topic", label, log)

        track = track_by_id.get(track_id)
        if not track:
            log.error(f"{label}: Track_ID {track_id!r} does not exist in Tracks sheet")
            continue

        subject = subject_by_id.get(subject_id)
        if not subject:
            log.error(f"{label}: Subject_ID {subject_id!r} does not exist in Subjects sheet")
            continue
        if subject.get("track_id") != track_id:
            log.error(
                f"{label}: Subject_ID {subject_id!r} belongs to Track_ID {subject.get('track_id')!r}, "
                f"not {track_id!r}"
            )
            continue

        module = module_by_id.get(module_id)
        if not module:
            log.error(f"{label}: Module_ID {module_id!r} does not exist in Modules sheet")
            continue
        if module.get("track_id") != track_id or module.get("subject_id") != subject_id:
            log.error(
                f"{label}: Module_ID {module_id!r} belongs to Track_ID {module.get('track_id')!r} "
                f"and Subject_ID {module.get('subject_id')!r}, not {track_id!r}/{subject_id!r}"
            )
            continue
        track_title = track["title"]
        subject_title = subject["title"]
        module_title = module["title"]

        difficulty = parse_difficulty(require_value(row, "Difficulty", label, log), label, log)
        knowledge_layer = parse_knowledge_layer(require_value(row, "Knowledge_Layer", label, log), label, log)
        content_status = parse_content_status(require_value(row, "Content_Status", label, log), label, log)
        course_status = parse_availability_status(row.get("Course_Status", ""), "Course_Status", label, log)
        diagnostic_status = parse_availability_status(row.get("Diagnostic_Status", ""), "Diagnostic_Status", label, log)
        learning_time = parse_learning_time(row.get("Learning_Time", ""), label, log)
        lens_relevance = parse_lens_relevance(row.get("Lens_Relevance", ""), label, log)
        radar_start = parse_date(row.get("Radar_Start_Date", ""), label, "Radar_Start_Date", log)
        radar_end = parse_date(row.get("Radar_End_Date", ""), label, "Radar_End_Date", log)
        validate_radar_dates(radar_start, radar_end, label, log)
        is_radar = parse_bool(row.get("Is_Radar", ""), label, log)
        radar_classification = parse_radar_classification(
            row.get("Radar_Classification", ""),
            is_radar,
            label,
            log,
        )

        if is_radar and not radar_start:
            log.warn(f"{label}: Is_Radar is TRUE but Radar_Start_Date is blank")
        if not row.get("Summary", "").strip():
            log.warn(f"{label}: Summary is blank")
        if not row.get("Why_It_Matters", "").strip():
            log.warn(f"{label}: Why_It_Matters is blank")
        if content_status == "Archived" and is_radar:
            log.warn(f"{label}: Archived topic is still marked Is_Radar")

        previous_module = previous_locations.get(topic_id)
        if previous_module and previous_module != module["slug"]:
            log.warn(f"{label}: topic moved from module {previous_module!r} to {module['slug']!r}")

        topic_name_by_module[(module["id"], topic_title)] += 1
        if topic_name_by_module[(module["id"], topic_title)] > 1:
            log.warn(f"{label}: duplicate topic name {topic_title!r} within module {module_title!r}")

        topic = {
            "id": topic_id,
            "topic_id": topic_id,
            "title": topic_title,
            "track_title": track_title,
            "subject_title": subject_title,
            "module_title": module_title,
            "track_id": track["id"],
            "subject_id": subject["id"],
            "module_id": module["id"],
            "slug": slugify(topic_title),
            "layer": knowledge_layer,
            "knowledge_layer": knowledge_layer,
            "difficulty": difficulty,
            "learning_time": learning_time,
            "lens_relevance": lens_relevance,
            "roles": list(lens_relevance.keys()),
            "status": STATUS_TO_JSON[content_status],
            "content_status": content_status,
            "priority": infer_priority(lens_relevance),
            "is_radar": is_radar,
            "radar_week": date_to_iso_week(radar_start),
            "radar_start_date": radar_start,
            "radar_end_date": radar_end,
            "radar_classification": radar_classification,
            "summary": row.get("Summary", "").strip(),
            "why_it_matters": row.get("Why_It_Matters", "").strip(),
            "resources": parse_resources(row.get("Resources", ""), label, log),
            "order": len(topics_by_module[module["id"]]) + 1,
            "course_status": course_status,
            "diagnostic_status": diagnostic_status,
            "hasCourse": course_status == "yes",
        }
        topics_by_module[module["id"]].append(topic)

    if log.errors:
        return {"tracks": []}

    modules_by_subject: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for module in sorted(module_by_id.values(), key=lambda item: item["order"]):
        module_topics = topics_by_module.get(module["id"], [])
        module["topics"] = module_topics
        module.update(aggregate_topics(module_topics))
        modules_by_subject[module["subject_id"]].append(module)

    subjects_by_track: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for subject in sorted(subject_by_id.values(), key=lambda item: item["order"]):
        subject_modules = modules_by_subject.get(subject["id"], [])
        subject["modules"] = subject_modules
        subject.update(aggregate_topics([topic for module in subject_modules for topic in module["topics"]]))
        subjects_by_track[subject["track_id"]].append(subject)

    tracks: list[dict[str, Any]] = []
    for track in sorted(track_by_id.values(), key=lambda item: item["order"]):
        track_subjects = subjects_by_track.get(track["id"], [])
        track["subjects"] = track_subjects
        track.update(
            aggregate_topics(
                [
                    topic
                    for subject in track_subjects
                    for module in subject["modules"]
                    for topic in module["topics"]
                ]
            )
        )
        tracks.append(track)

    return {"tracks": tracks}


def parse_hierarchy_row(row: dict[str, str], node_type: str, node_id: str, title: str, label: str, log: BuildIssueLog) -> dict[str, Any]:
    radar_start = parse_date(row.get("Radar_Start_Date", ""), label, "Radar_Start_Date", log)
    radar_end = parse_date(row.get("Radar_End_Date", ""), label, "Radar_End_Date", log)
    validate_radar_dates(radar_start, radar_end, label, log)
    content_status = parse_content_status(row.get("Content_Status", "Published").strip() or "Published", label, log)
    course_ids = parse_course_ids(row.get("Course_ID(s)", ""))
    return {
        "id": node_id,
        "title": title,
        "slug": slugify(title),
        "order": parse_order(row.get("Order", ""), label, log),
        "summary": row.get("Summary", "").strip(),
        "course_ids": course_ids,
        "hasCourse": bool(course_ids),
        "knowledge_layer": row.get("Knowledge_Layer", "").strip(),
        "why_it_matters": row.get("Why_It_Matters", "").strip(),
        "resources": parse_resources(row.get("Resources", ""), label, log),
        "is_radar": parse_bool(row.get("Is_Radar", ""), label, log),
        "radar_start_date": radar_start,
        "radar_end_date": radar_end,
        "content_status": content_status,
        "status": STATUS_TO_JSON[content_status],
        "node_type": node_type,
    }


def validate_radar_dates(start: str | None, end: str | None, label: str, log: BuildIssueLog) -> None:
    if start and end and end < start:
        log.error(f"{label}: Radar_End_Date cannot be earlier than Radar_Start_Date")


def infer_priority(lens_relevance: dict[str, str]) -> str:
    if "AI Engineer" in lens_relevance:
        return RELEVANCE_TO_PRIORITY[lens_relevance["AI Engineer"]]
    if "Must" in lens_relevance.values():
        return "high"
    if "Good" in lens_relevance.values():
        return "medium"
    return "low"


def aggregate_topics(topics: list[dict[str, Any]]) -> dict[str, Any]:
    difficulty_counts = Counter(topic["difficulty"] for topic in topics)
    knowledge_layer_counts = Counter(topic["knowledge_layer"] for topic in topics if topic.get("knowledge_layer"))
    lens_counts: dict[str, Counter[str]] = defaultdict(Counter)
    learning_time = 0.0
    has_learning_time = False
    for topic in topics:
        if topic.get("learning_time") is not None:
            learning_time += float(topic["learning_time"])
            has_learning_time = True
        for role, category in topic.get("lens_relevance", {}).items():
            lens_counts[role][category] += 1
    return {
        "difficulty": {key: difficulty_counts.get(key, 0) for key in sorted(VALID_DIFFICULTIES)},
        "knowledge_layer": dict(sorted(knowledge_layer_counts.items())),
        "learning_time": round(learning_time, 2) if has_learning_time else None,
        "lens_relevance": {
            role: {category: counts.get(category, 0) for category in ["Must", "Good", "Optional"]}
            for role, counts in sorted(lens_counts.items())
        },
    }


def load_previous_topic_locations() -> dict[str, str]:
    if not PREVIOUS_JSON_FILE.exists():
        return {}
    try:
        data = json.loads(PREVIOUS_JSON_FILE.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return {}
    locations: dict[str, str] = {}
    for track in data.get("tracks", []):
        for subject in track.get("subjects", []):
            for module in subject.get("modules", []):
                for topic in module.get("topics", []):
                    locations[topic.get("id", "")] = module.get("slug", "")
    return locations


def count_nodes(syllabus: dict[str, Any]) -> tuple[int, int, int, int]:
    tracks = len(syllabus["tracks"])
    subjects = sum(len(track["subjects"]) for track in syllabus["tracks"])
    modules = sum(len(subject["modules"]) for track in syllabus["tracks"] for subject in track["subjects"])
    topics = sum(
        len(module["topics"])
        for track in syllabus["tracks"]
        for subject in track["subjects"]
        for module in subject["modules"]
    )
    return tracks, subjects, modules, topics


def update_workbook_aggregate_columns(path: Path, syllabus: dict[str, Any]) -> None:
    """Write derived parent metrics back into the workbook.

    Authors should not maintain these cells manually. They are refreshed from
    Topic rows every time the build succeeds.
    """
    wb = load_workbook(path)

    track_nodes = {track["id"]: track for track in syllabus["tracks"]}
    subject_nodes = {
        subject["id"]: subject
        for track in syllabus["tracks"]
        for subject in track["subjects"]
    }
    module_nodes = {
        module["id"]: module
        for track in syllabus["tracks"]
        for subject in track["subjects"]
        for module in subject["modules"]
    }

    update_sheet_aggregates(wb["Tracks"], "Track_ID", track_nodes)
    update_sheet_aggregates(wb["Subjects"], "Subject_ID", subject_nodes)
    update_sheet_aggregates(wb["Modules"], "Module_ID", module_nodes)
    wb.save(path)


def update_sheet_aggregates(ws: Any, id_column: str, nodes_by_id: dict[str, dict[str, Any]]) -> None:
    header = {cell.value: index for index, cell in enumerate(ws[1], start=1)}
    required = {"Knowledge_Layer", "Difficulty", "Learning_Time", "Lens_Relevance", id_column}
    if not required.issubset(header):
        return

    for row_index in range(2, ws.max_row + 1):
        node_id = cell_value(ws.cell(row_index, header[id_column]).value)
        node = nodes_by_id.get(node_id)
        if not node:
            continue
        ws.cell(row_index, header["Knowledge_Layer"]).value = json.dumps(node.get("knowledge_layer", {}), sort_keys=True)
        ws.cell(row_index, header["Difficulty"]).value = json.dumps(node.get("difficulty", {}), sort_keys=True)
        learning_time = node.get("learning_time")
        ws.cell(row_index, header["Learning_Time"]).value = f"{learning_time or 0}h"
        ws.cell(row_index, header["Lens_Relevance"]).value = json.dumps(node.get("lens_relevance", {}), sort_keys=True)


def write_report(log: BuildIssueLog, syllabus: dict[str, Any] | None = None, path: Path = BUILD_REPORT_FILE) -> None:
    tracks = subjects = modules = topics = 0
    if syllabus:
        tracks, subjects, modules, topics = count_nodes(syllabus)

    title = "Syllabus Migration Report" if path == REPORT_FILE else "Syllabus Build Report"
    lines = [
        f"# {title}",
        "",
        f"- Migrated legacy workbook: {'Yes' if log.migrated else 'No'}",
        f"- Tracks: {tracks}",
        f"- Subjects: {subjects}",
        f"- Modules: {modules}",
        f"- Topics: {topics}",
        "",
        "## Records Migrated",
    ]
    if log.migrated_records:
        lines.extend(f"- {key}: {value}" for key, value in log.migrated_records.items())
    else:
        lines.append("- No legacy migration was required.")

    lines.extend(["", "## Manual Review"])
    lines.extend(f"- {item}" for item in log.manual_review) if log.manual_review else lines.append("- None")
    lines.extend(["", "## Warnings"])
    lines.extend(f"- {item}" for item in log.warnings) if log.warnings else lines.append("- None")
    lines.extend(["", "## Errors"])
    lines.extend(f"- {item}" for item in log.errors) if log.errors else lines.append("- None")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> int:
    log = BuildIssueLog()
    try:
        if not INPUT_FILE.exists():
            raise FileNotFoundError(f"Input file not found: {INPUT_FILE}")

        if not has_target_workbook(INPUT_FILE):
            migrate_legacy_workbook(INPUT_FILE, log)

        ensure_topic_authoring_columns(INPUT_FILE)

        track_rows = load_workbook_rows(INPUT_FILE, "Tracks", TRACK_COLUMNS)
        subject_rows = load_workbook_rows(INPUT_FILE, "Subjects", SUBJECT_COLUMNS)
        module_rows = load_workbook_rows(INPUT_FILE, "Modules", MODULE_COLUMNS)
        topic_rows = load_workbook_rows(INPUT_FILE, "Topics", TOPIC_COLUMNS)
        syllabus = build_syllabus(track_rows, subject_rows, module_rows, topic_rows, log)

        if log.errors:
            write_report(log, syllabus, BUILD_REPORT_FILE)
            print(f"Error: validation failed; see {BUILD_REPORT_FILE}", file=sys.stderr)
            for error in log.errors:
                print(f"- {error}", file=sys.stderr)
            return 1

        update_workbook_aggregate_columns(INPUT_FILE, syllabus)
        output_json = json.dumps(syllabus, indent=2) + "\n"
        OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
        SRC_OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
        OUTPUT_FILE.write_text(output_json, encoding="utf-8")
        SRC_OUTPUT_FILE.write_text(output_json, encoding="utf-8")
        write_report(log, syllabus, BUILD_REPORT_FILE)
        if log.migrated:
            write_report(log, syllabus, REPORT_FILE)

        tracks, subjects, modules, topics = count_nodes(syllabus)
        print(
            f"Done: {tracks} tracks, {subjects} subjects, {modules} modules, "
            f"{topics} topics exported to {OUTPUT_FILE} and {SRC_OUTPUT_FILE}"
        )
        if log.warnings:
            print(f"Warnings: {len(log.warnings)}; see {BUILD_REPORT_FILE}")
        if log.manual_review:
            print(f"Manual review items: {len(log.manual_review)}; see {REPORT_FILE}")
        return 0
    except Exception as exc:
        log.error(str(exc))
        write_report(log, path=BUILD_REPORT_FILE)
        print(f"Error: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
